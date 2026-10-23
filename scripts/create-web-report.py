"""生成本次联网实现的代码功能书；读取源码提取真实行号，不读取环境密钥。"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
import re

root = Path(__file__).resolve().parents[1]
out = root.parent / '智游闽韵_RAG联网代码功能书.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.8)
sec.left_margin = sec.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Microsoft YaHei'
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
doc.add_heading('智游闽韵 RAG联网代码功能书', 0)
doc.add_paragraph('实现快照 2026年10月17日')
doc.add_paragraph('适用项目：' + str(root))
doc.add_paragraph('本次在原有本地RAG之后增加百炼联网补充、可点击网页来源与失败降级。网络结果不写入向量知识库。本文解释本次新增和修改的逻辑块，并给出已有模块在完整链路中的位置；不是对整个项目每一行代码的重复抄录。行号对应生成本报告时的源码，后续编辑会改变行号。')

def h(t): doc.add_heading(t, 1)
def p(t): doc.add_paragraph(t)
def block(file, start, end, title, explanation):
    lines = (root / file).read_text(encoding='utf-8-sig').splitlines()
    a = next(i for i, s in enumerate(lines) if start in s)
    b = next((i for i in range(a+1, len(lines)) if end and end in lines[i]), len(lines))
    doc.add_heading(title, 2)
    p(str(root / file) + f'\n第 {a+1} 至 {b} 行')
    p(explanation)
    # 长提示词按字符换行，保证Word中不会越过页边界；保留源码行号。
    for i in range(a, b):
        line = f'{i+1:03d}  {lines[i]}'
        chunks = [line[n:n+75] for n in range(0, len(line), 75)] or ['']
        for part in chunks:
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1
            run = para.add_run(part)
            run.font.name = 'Consolas'; run.font.size = Pt(8)
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

h('一 本次完成的功能与边界')
p('本地资料充分时仍使用本地RAG，不增加搜索开销。本地未命中、证据不足、局部缺失或本地回答链路异常时，最多尝试一次联网。开放时间、天气等时效问题优先联网；问答不代替预约或退款操作。')
p('联网使用当前服务端RAG密钥，通过百炼原生接口调用qwen-plus。只有配置可识别的官方域名、有效密钥、账号额度与服务权限时才能成功。现有密钥实测返回了8个引用网页，约16.091秒；其中包含福建省文物局页面，也包含非官方网页。“优先官网”是模型提示，不是强制官网白名单。')
p('来源验证检查引用编号、标题、URL和响应结束状态，不能证明网页内容全部正确，也没有独立抓取全文核实事实。外链可能后续失效。accessedAt为本次查询时间，不是网页发布时间。暂不自动入库、暂不实现自建爬虫或独立网络结果重排。')

h('二 从提问到回答的完整流程')
for text in [
'1 用户在AskExperience提交主问题、语言和会话ID。前端以SSE接收状态及正文；没有重新引入三个子问题拆解。',
'2 /api/chat/answer校验来源、请求大小、登录与会话归属，并限流。读取最近6条历史消息，将请求ID贯穿日志。',
'3 pipeline区分寒暄、知识问题与实时需求。context还原明确指代；存在多个对象则先澄清。问答不执行订单动作。',
'4 实时问题进入联网分支。其他知识问题从.data/rag/index.json载入本地索引，经BM25与稠密向量检索、元数据过滤、RRF融合、可用时重排、去重和父块扩展得到证据。向量模型不匹配或服务失败时保留已有词法降级行为。',
'5 模型判断证据覆盖情况。充分时生成声明并进行引文、数字与语义检查；部分覆盖时先核验能回答的内容，再联网补充missing部分；没有证据时直接尝试联网。',
'6 web.ts调用百炼原生联网接口，由供应商完成搜索和生成。SSE增量正文先作为临时内容展示，收齐结果后核对实际引用元数据。该分支不是把网页切成向量再走本地重排。',
'7 联网成功返回web模式以及webSources；部分覆盖保留本地sections和sources。失败撤回临时网络正文：普通知识问题降级为无来源的模型回答，部分覆盖保留已核验内容和有标注的建议，实时问题明确无法确认。',
'8 生成最多3个推荐追问，记录回答耗时与首次正文时间。登录用户将正文、本地及网页来源保存为会话JSON；游客仅当前页面可见。前端正文上方显示网页数量，点击打开来源弹窗，再点击网页标题访问外部网页。',
'9 网络问答只写会话历史，不调用buildIndex；知识库导入仍由rag:ingest独立完成。']:
    p(text)

h('三 联网适配器逐段解释')
block('features/rag/web.ts','type SearchRow','/** 只允许','供应商数据结构','SearchRow只接收编号、标题与URL。NativeResponse描述百炼原生SSE返回的choices和search_info；unknown要求引用在使用前校验，不能直接把模型字符串当链接。')
block('features/rag/web.ts','export function nativeSearchEndpoint','/** 可用','官方端点转换','先用URL解析配置，拒绝HTTP、账号密码和非标准端口，再通过域名白名单构造原生接口。普通OpenAI兼容接口不提供这里需要的search_info，因此不是简单给旧接口加一个参数。')
block('features/rag/web.ts','export function webReady','/** 移除','本地可用性判断','检查开关、密钥和官方端点是否齐全。不做额度或权限探测，真正可用性取决于远端调用。')
block('features/rag/web.ts','export function cleanWebText','/** 校验','正文清理','移除ref角标和自由URL，正文不再显示供应商原始引用格式。真正可点击链接仅从经过校验的webSources生成，避免模型随意写链接进入页面。')
block('features/rag/web.ts','export function validateWebResult','/** 一次','引用验证与去重','要求finish_reason为stop，拒绝空正文、超长正文和无引用答案。提取每个ref编号，对应真实返回的search_results，检查标题与URL。重复网页仅计算一次。这里证明编号和网页元数据一致，不证明每句事实均受网页支持。')
block('features/rag/web.ts','export async function searchAnswer','    if (!response.ok)','请求参数与提示词','读取配置，合并用户取消信号和超时。Authorization只在服务端请求中使用；redirect:error防止带凭据重定向。提示词要求回答主问题、遵循语言和分段排版，并提醒忽略资料指令。forced_search要求搜索；enable_source和enable_citation要求返回网页及引用；prepend_search_result使流中先收到来源；incremental_output启用增量正文。')
block('features/rag/web.ts','    if (!response.ok)',None,'流解析与完成条件','非2xx或非SSE立即失败。readSse负责跨字节和帧边界；本函数汇总来源和正文，限制响应尺寸与来源数。onDraft通知前端临时正文。全部结束后再次检查取消、核对引用，成功记录web.sources.validated。失败由主流程撤回草稿。')
block('features/rag/web-url.ts','export function safeWebUrl',None,'网页链接过滤','客户端与服务端共用纯URL检查。拒绝危险协议、用户信息、IP字面量和常见内网后缀。它不是DNS级SSRF防护；本项目不在后端请求这些网页，用户点击后由浏览器访问。')

h('四 主流程修改解释')
block('features/rag/pipeline.ts','  async function tryWeb','  /**','联网统一入口','webTried限制每个主问题最多一次联网；trace记录触发原因和结果。partial传入已验证内容，join把它与联网补充拼接。来源核对通过才构造最终web回答；任何异常先replace为原本地内容，防止失败网络正文继续留在页面。用户取消则抛出而非再降级发新请求。')
block('features/rag/pipeline.ts','  async function modelAnswer','  const useModel','一般知识的失败降级','先尝试联网，成功直接返回。联网不可用才走原有普通模型模式，明确不附虚构引用。它与实时问题的保守分支不同，不能把模型一般知识包装为今日公告。')
block('features/rag/pipeline.ts','  // 问答不执行订单动作','  // 先还原','实时信息与订单边界','识别预约下单退款动作并引导票务页面；实时查询尝试联网。失败时返回无法确认，而不是调用静态知识库给出当前状态。识别使用关键词，存在漏判和误判风险，后续需通过真实问题集持续调整。')
block('features/rag/pipeline.ts','    if (audit.status === "partial")','    // Never stream','部分覆盖保留本地证据','先将已经验证的claim转换为本地引用和正文，再把已回答内容与缺失需求传给联网适配器。联网失败时继续原有补充建议逻辑；充分覆盖则记录跳过联网。')
p('其他接入点：索引加载失败、零命中、审查none、声明修正失败和主链异常均经modelAnswer尝试联网；核心证据冲突可联网补充，失败仍拒绝确定性结论。已在进行的旧兼容模型调用尚未全部接入取消信号，取消主要能中止新的联网请求并阻止后续保存。')

h('五 数据结构与前端修改')
block('features/rag/types.ts','export type WebSource','/** 多轮','回答契约','WebSource与本地Citation分开，避免给网页伪造行号和TXT片段。mode新增web；webSources和trace.web都是可选字段，兼容旧会话JSON。无需数据库迁移。')
block('features/rag/config.ts','    // 联网独立模型','    embedding:','配置读取','默认启用联网，默认模型qwen-plus。超时默认30秒，限制在5至60秒。RAG_API_KEY优先于DASHSCOPE_API_KEY，base配置必须是支持的官方域名；修改配置后重启开发服务。')
block('features/ai-chat/WebSources.tsx','export function WebSources',None,'网页来源组件','渲染时再次校验URL，React自动转义标题。使用noopener noreferrer隔离外部窗口；域名独占一行，用户可以辨认来源站点。没有把标题HTML直接插入页面。')
block('features/ai-chat/AskExperience.tsx','              {!!turn.answer.webSources','              <footer>','正文上方网页入口','按最终去重后的webSources显示网页数量，点击打开既有来源弹窗。它和本地查询数分开，不假装知道百炼内部搜索了多少关键词。正文仍用AnswerContent处理加粗与分段。')
block('features/ai-chat/AskExperience.tsx','      {!!sourceTurn?.answer.webSources','    </dialog>','合并来源弹窗','弹窗上部显示网页链接，下部保留本地引用预览。仅有网页时不显示本地TXT提示。沿用全宽对话，不恢复旧右侧来源栏。')
block('features/rag/followups.ts','export async function generateFollowUps',None,'推荐追问兼容联网','纯网页答案没有本地sections，因此取intro作为推荐问题的上下文。局部回答仍附带advice。推荐问题失败不会破坏正文；最多3条，点击后正常发起下一轮。')

h('六 接口保存与取消')
block('app/api/chat/answer/route.ts','  // 取消浏览器响应','  /**','请求取消信号','把浏览器请求中止与SSE取消控制器合并，传给answerQuestion。客户端断开后会尝试终止联网HTTP读取，但供应商已产生的费用不保证退回。')
block('app/api/chat/answer/route.ts','    const answer = await answerQuestion','    let persisted','问答执行与推荐追问','传入requestSignal，问答完成和推荐追问完成后都检查取消状态。已有计时与推荐逻辑保留。历史保存前不再继续已取消的请求。')
block('app/api/chat/answer/route.ts','          await tx.chatMessage.create({ data: { sessionId: session.id, role: "assistant"','          return session.id','会话持久化','完整answer写content JSON；citations增加webSources。写的是用户会话记录，不是向量索引，也没有把网页来源伪装成本地文档。')
p('同文件SSE的cancel回调设置disconnected并调用cancellation.abort。finally仍释放限流额度。普通JSON客户端也共用execute逻辑。')

h('七 本地知识库和既有模块')
for file, text in [
('features/rag/ingest.ts','parseDocument按标题或问答形成父子块；buildIndex生成向量并以临时文件替换完整索引。网页回答不会调用这里。'),
('features/rag/store.ts','loadIndex读取索引，按文件状态缓存并共享并发加载，记录版本、块数与向量维度。'),
('features/rag/retrieve.ts','tokenize、bm25、cosine、fuse、retrieve分别负责分词、词法评分、向量相似度、RRF融合与检索编排；按配置调用重排。'),
('features/rag/context.ts','normalizeHistory整理历史；resolveReference还原指代，保留单个主问题。'),
('features/rag/providers.ts','旧兼容接口负责普通聊天、Embedding与重排。与新增web.ts并存，不把原有RAG替换为纯搜索。'),
('features/rag/stream.ts','解析SSE并定义status、delta、replace事件。replace可替换或清空临时正文。'),
('features/rag/logger.ts','为请求绑定追踪ID并输出结构化日志，避免记录API密钥。'),
('features/rag/limits.ts','现有限流为单进程限制；多实例生产部署需要共享限流，不等同于全局配额。')]:
    p(str(root/file)+'\n'+text)
p('本链路默认向量存储：'+str(root/'.data/rag/index.json')+'。这是带向量数组的本地JSON索引，不是本次新建的Qdrant集合。配置RAG_DATA_DIR可改目录；RAG_SOURCE_DIR控制TXT来源目录。迁移电脑时需携带索引和原文，保持向量模型及维度一致。')

h('八 新增日志与排查')
for t in ['web.decision：说明本次是否具备联网配置及触发或跳过原因。', 'web.search.start：开始调用联网模型，记录模型名。', 'web.search.complete：联网适配器成功返回，记录调用耗时。', 'web.search.failed：调用、流解析或引用验证失败；安全错误码供定位。', 'web.sources.validated：返回来源数量returned与实际引用去重数量cited，不是事实真实性评分。', 'web.fallback：网络草稿被撤回并开始降级。', 'pipeline.stage：真实进入的前端阶段；联网分支没有自行执行本地向量重排，不会伪造重排阶段。']:
    p(t)
p('WEB_HTTP_401或403需检查密钥、地域和权限；429表示限流或配额限制；WEB_NO_CITATIONS表示没有实际引用；WEB_INVALID_CITATION表示编号未对应安全有效的来源；WEB_INCOMPLETE_RESPONSE表示输出未正常结束；WEB_INVALID_STREAM表示响应不是预期的SSE。日志不打印认证头或远端原始错误正文。')

h('九 测试与运行方式')
p('自动化回归：npm run rag:test。当前53项通过，其中12项为新联网测试，41项为既有RAG测试。修改旧实时用例断言，使其在无可用联网时拒绝猜测而非转普通模型。测试模拟网络，不代表所有线上问题都通过。TypeScript检查与本次修改文件ESLint检查通过。')
p('真实适配器测试：npm run rag:web:smoke。scripts/rag-web-smoke.ts加载本地环境并发起一次付费可能的搜索，输出耗时和来源，不打印密钥、不写索引。当前密钥已通过一次真实调用；未宣称完整浏览器交互验收已完成。')
p('新增测试文件：'+str(root/'tests/rag-web.test.ts')+'。覆盖官方端点、危险链接、引用映射与去重、SSE参数与草稿、实时成功、429失败、无效引用撤回、一般问题降级、关闭及取消、来源组件安全HTML。')
p('配置示例在'+str(root/'.env.example')+'：RAG_WEB_ENABLED=true，RAG_WEB_MODEL=qwen-plus，RAG_WEB_TIMEOUT_MS=30000。服务端使用现有RAG_API_KEY或DASHSCOPE_API_KEY。分享项目时不要打包.env和.env.local，让其他成员配置自己的密钥与账号权限。')
p('启动：在项目目录运行npm run dev。可测试“今天福建博物院开放吗”，检查正在联网搜索、网页数量、来源弹窗；再把RAG_WEB_ENABLED设为false并重启，测试实时问题是否明确无法确认。普通知识资料不足时可降级为有提示的模型回答。')
p('回退联网行为只需RAG_WEB_ENABLED=false并重启，不需要删除索引或撤销其他业务代码。新增WebSource为可选字段，既有历史仍可读取。不要对当前包含其他修改的仓库执行整体重置。')

h('十 限制和后续验收')
p('建议后续用真实问答集检验网页回答的准确性、英语输出、最新公告日期、费用和超时比例。当前只做引用对应关系校验；搜索可能包含非官网、旧公告和相互矛盾页面，不能保证余票和开放状态实时准确。浏览器来源链接是否能打开，还受对方网站访问策略影响。')
p('本次没有重新执行100题真实付费评测，没有增加自动入库，没有实现网络全文级证据绑定，也没有把所有兼容模型请求改为可取消。首次正文计时包含来源尚待核对的临时文本；如果草稿被撤回，该时间不应解释为最终可靠正文时间。文档已检查可读取结构，但当前环境缺少LibreOffice，无法完成逐页渲染检查。')
p('官方接口依据：https://help.aliyun.com/zh/model-studio/web-search 。本次按原生DashScope的搜索参数和search_info实现；模型支持与计费以当前账号控制台为准。')
doc.save(out)
print(str(out))
print('paragraphs', len(doc.paragraphs))
